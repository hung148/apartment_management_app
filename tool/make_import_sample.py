"""Made-up file in the old app's export format, for testing the sheet import
on staging (2026-10-05). Every name, phone and amount here is invented; the
real customer file is never used for tests (IMPORT_SHEET.md).

    python tool/make_import_sample.py            -> tool/import_sample.xlsx

Dates are relative to today, so the file always has past, current and future
stays. It also carries what the import must ignore or clean up: a password
column with fake values, the unnamed change-log column, junk phone numbers,
a phone number that lost its leading 0, 0 đ payments, overlapping stays and a
room marked "Thuê dài hạn" with no lease.
"""
import datetime as dt
import json
import os

import openpyxl

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'import_sample.xlsx')
today = dt.datetime.combine(dt.date.today(), dt.time())


def at(days, hour=14, minute=0):
    return today + dt.timedelta(days=days, hours=hour, minutes=minute)


def pays(*items):
    return json.dumps([{'ngayThanhToan': d.strftime('%Y-%m-%d'), 'soTien': a, 'ghiChu': n} for d, a, n in items],
                      ensure_ascii=False)


wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Đặt phòng'
ws.append(['ID', 'Tên cơ sở', 'Mã phòng', 'Ngày tạo', 'Tên khách hàng', 'Số điện thoại', 'Nguồn', 'Nhân viên',
           'Hình thức phòng', 'Ngày giờ checkin', 'Ngày giờ checkout', 'Đơn giá', 'Phụ phí', 'Tổng doanh thu',
           'Ghi chú', 'Thanh toán Json', 'Số tiền còn lại', 'Dọn phòng Json', None])
log = json.dumps([{'t': '01/10/2026 10:00', 'u': 'Mẫu', 'a': 'A', 'id': 'X1', 'd': 'Thêm mới'}], ensure_ascii=False)
rows = [
    # id, building, room, guest, phone, source, staff, method, in, out, price, fee, payments, note
    ('MK01', 'Nhà Mẫu A', 'P101', 'Khách Mẫu 01', 'ko có', 'Chính chủ', 'Lan Mẫu', '2N1D', at(-20), at(-19, 12), 500000, 0,
     [(at(-25), 200000, 'cọc'), (at(-19), 300000, '')], ''),
    ('MK02', 'Nhà Mẫu A', 'P101', 'Khách Mẫu 02', 912000001.0, 'Booking', 'Lan Mẫu', '3N2D', at(-10), at(-8, 12), 900000, 100000,
     [(at(-8), 1000000, '')], 'Thêm nệm'),
    # Overlaps MK02 by one night.
    ('MK03', 'Nhà Mẫu A', 'P101', 'Khách Mẫu 03', 'xxx', 'Cộng tác viên', 'Hùng Mẫu', 'Qua đêm', at(-9, 21), at(-8, 10), 350000, 0,
     [(at(-9), 350000, '')], ''),
    ('MK04', 'Nhà Mẫu A', 'P102', 'Khách Mẫu 04', '0901000004', 'Facebook Ads', 'Hùng Mẫu', 'Thuê giờ', at(-3, 10), at(-3, 13), 210000, 0,
     [(at(-3), 0, ''), (at(-3), 210000, '')], ''),
    # Staying now, deposit only.
    ('MK05', 'Nhà Mẫu A', 'P102', 'Khách Mẫu 05', 'k', 'Chính chủ', 'Lan Mẫu', '4N3D', at(-1), at(2, 12), 1500000, 0,
     [(at(-4), 500000, 'tiền cọc')], ''),
    # Future stays.
    ('MK06', 'Nhà Mẫu A', 'P103', 'Khách Mẫu 06', '0901 000 006', 'Giới thiệu', 'Lan Mẫu', '2N1D', at(5), at(6, 12), 550000, 0,
     [(at(-2), 200000, 'cọc')], ''),
    ('MK07', 'Nhà Mẫu B', 'P201', 'Khách Mẫu 07', 'Ko có', 'Booking', 'Hùng Mẫu', 'Combo 2N1D', at(8), at(9, 12), 700000, 0, [], ''),
    # Long stays: labelled, running now; unlabelled 25 days, ended.
    ('MK08', 'Nhà Mẫu B', 'P202', 'Người Thuê Mẫu 08', '0901000008', None, 'Lan Mẫu', 'Thuê dài hạn', at(-40), at(50, 12), 15000000, 0,
     [(at(-40), 5000000, 'tiền cọc'), (at(-40), 5000000, 'tiền thuê tháng 1'), (at(-10), 5000000, 'tiền thuê tháng 2')], ''),
    ('MK09', 'Nhà Mẫu B', 'P203', 'Người Thuê Mẫu 09', 'không có', None, 'Hùng Mẫu', None, at(-60), at(-35, 12), 6000000, 0,
     [(at(-60), 6000000, '')], ''),
    # Room missing from the Phòng tab; overpaid.
    ('MK10', 'Nhà Mẫu B', 'P299', 'Khách Mẫu 10', '0901000010', 'Chính chủ', 'Ai Đó', 'Qua đêm', at(-5, 20), at(-4, 9), 300000, 0,
     [(at(-4), 350000, '')], ''),
]
for (i, b, r, g, ph, src, staff, method, cin, cout, price, fee, plist, note) in rows:
    paid = sum(a for _, a, _ in plist)
    ws.append([i, b, r, cin - dt.timedelta(days=3), g, ph, src, staff, method, cin, cout, price, fee, price + fee, note,
               pays(*plist), price + fee - paid, '{}', log])
for _ in range(5):
    ws.append([None] * 19)

ws = wb.create_sheet('Phòng')
ws.append(['ID', 'Tên cơ sở', 'Giá thuê cơ sở', 'Mã phòng Json', 'Ghi chú'])
ws.append(['MB01', 'Nhà Mẫu A', 25000000.0, json.dumps([{'maPhong': p, 'trangThai': 'Đã dọn', 'ghiChu': ''} for p in ['P101', 'P102', 'P103']]), 'Tòa mẫu A'])
ws.append(['MB02', 'Nhà Mẫu B', 0.0, json.dumps([{'maPhong': 'P201', 'trangThai': 'Đã dọn', 'ghiChu': ''},
                                                 {'maPhong': 'P202', 'trangThai': 'Thuê dài hạn', 'ghiChu': ''},
                                                 {'maPhong': 'P203', 'trangThai': 'Đã dọn', 'ghiChu': ''},
                                                 {'maPhong': 'P204', 'trangThai': 'Thuê dài hạn', 'ghiChu': 'không có hợp đồng trong file'}],
                                                ensure_ascii=False), 'Tòa mẫu B'])

ws = wb.create_sheet('Nhân viên')
ws.append(['ID', 'Tên nhân viên', 'Tên đăng nhập', 'Mật khẩu', 'Quyền', 'Quản lý cơ sở', 'Màu', 'Tỷ lệ hoa hồng', 'Hoa hồng'])
ws.append(['MS01', 'Chủ Mẫu', 'chumau', 'mat-khau-gia-1', 'Admin', '[]', '#2c5aa0', 0.0, 0])
ws.append(['MS02', 'Lan Mẫu', 'lanmau', 'mat-khau-gia-2', 'Quản lý', '[]', '#24c6b3', 8.0, 0])
ws.append(['MS03', 'Hùng Mẫu', 'hungmau', 123456.0, 'Nhân viên', '[]', '#f9ab2f', 0.0, 0])

ws = wb.create_sheet('Chi phí')
ws.append(['ID', 'Ngày chi', 'Cơ sở', 'Nhóm chi phí', 'Nội dung chi', 'Số tiền', 'Ghi chú'])
ws.append(['ME01', at(-15, 0), 'Nhà Mẫu A', 'Chi phí thuê nhà', 'Tiền nhà tháng trước', 25000000.0, ''])
ws.append(['ME02', at(-6, 0), 'Nhà Mẫu B', 'Chi phí dọn phòng', 'Nước giặt', 180000.0, 'mua ở chợ'])

ws = wb.create_sheet('Cài đặt')
ws.append(['Key', 'Value', 'Ghi chú'])
ws.append(['sources', json.dumps(['Giới thiệu', 'Cộng tác viên', 'Booking', 'Chính chủ'], ensure_ascii=False), ''])
ws.append(['room_methods', json.dumps(['2N1D', '3N2D', '4N3D', 'Qua đêm', 'Thuê giờ', 'Thuê dài hạn'], ensure_ascii=False), ''])

ws = wb.create_sheet('Dọn phòng')
ws.append(['Lịch dọn phòng (bỏ qua khi nhập)'])

wb.save(OUT)
print('wrote', OUT)
